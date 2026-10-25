"""Render a local scene specification to editable PPTX and a same-layout PDF preview.

No experiment execution. Chinese font assets and internal content are supplied locally.
The PDF is generated from the same scene, not converted by Microsoft Office.
"""
from pathlib import Path
import json,math
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR,PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image

W,H=960,540
FONT='Microsoft YaHei'
def color(h):return tuple(int(h[i:i+2],16)/255 for i in (0,2,4))
def rgb(h):return RGBColor.from_string(h)
def register_fonts(regular,bold):
 pdfmetrics.registerFont(TTFont('BriefRegular',str(regular)))
 pdfmetrics.registerFont(TTFont('BriefBold',str(bold)))
def wrap(text,size,width,bold=False):
 font='BriefBold' if bold else 'BriefRegular';out=[]
 for paragraph in text.split('\n'):
  if not paragraph:out.append('');continue
  line=''
  for char in paragraph:
   if line and pdfmetrics.stringWidth(line+char,font,size)>width-4:out.append(line);line=char
   else:line+=char
  out.append(line)
 return out

def render(slides,pptx_path,pdf_path,regular,bold,workspace_root=None):
 root=Path(workspace_root or Path.cwd()).resolve()
 for candidate in [pptx_path,pdf_path,regular,bold]:assert Path(candidate).resolve().is_relative_to(root),candidate
 register_fonts(regular,bold);prs=Presentation();prs.slide_width=Pt(W);prs.slide_height=Pt(H)
 pdf=canvas.Canvas(str(pdf_path),pagesize=(W,H));pdf.setTitle('Research progress — same-layout preview');audit=[]
 for number,spec in enumerate(slides,1):
  slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=rgb(spec.get('bg','FFFFFF'))
  pdf.setFillColorRGB(*color(spec.get('bg','FFFFFF')));pdf.rect(0,0,W,H,fill=1,stroke=0)
  for item in spec['items']:
   kind=item['type'];x,y,w,h=[float(item.get(k,0)) for k in ['x','y','w','h']]
   assert x>=0 and y>=0 and x+w<=W+.1 and y+h<=H+.1,(number,item)
   if kind in ['rect','ellipse','chevron']:
    shape=slide.shapes.add_shape({'rect':MSO_SHAPE.ROUNDED_RECTANGLE if item.get('radius',0) else MSO_SHAPE.RECTANGLE,'ellipse':MSO_SHAPE.OVAL,'chevron':MSO_SHAPE.CHEVRON}[kind],Pt(x),Pt(y),Pt(w),Pt(h))
    fill=item.get('fill');stroke=item.get('stroke')
    if fill:shape.fill.solid();shape.fill.fore_color.rgb=rgb(fill)
    else:shape.fill.background()
    if stroke:shape.line.color.rgb=rgb(stroke);shape.line.width=Pt(item.get('line_width',1))
    else:shape.line.fill.background()
    if fill:pdf.setFillColorRGB(*color(fill))
    if stroke:pdf.setStrokeColorRGB(*color(stroke));pdf.setLineWidth(item.get('line_width',1))
    if kind=='rect':pdf.roundRect(x,H-y-h,w,h,item.get('radius',0),fill=int(bool(fill)),stroke=int(bool(stroke)))
    elif kind=='ellipse':pdf.ellipse(x,H-y-h,x+w,H-y,fill=int(bool(fill)),stroke=int(bool(stroke)))
    else:
     p=pdf.beginPath();p.moveTo(x,H-y);p.lineTo(x+w*.65,H-y);p.lineTo(x+w,H-y-h/2);p.lineTo(x+w*.65,H-y-h);p.lineTo(x,H-y-h);p.lineTo(x+w*.35,H-y-h/2);p.close();pdf.drawPath(p,fill=int(bool(fill)),stroke=int(bool(stroke)))
   elif kind=='line':
    shape=slide.shapes.add_connector(1,Pt(x),Pt(y),Pt(x+w),Pt(y+h));shape.line.color.rgb=rgb(item['color']);shape.line.width=Pt(item.get('width',1));pdf.setStrokeColorRGB(*color(item['color']));pdf.setLineWidth(item.get('width',1));pdf.line(x,H-y,x+w,H-y-h)
   elif kind=='image':
    path=Path(item['path']);assert path.resolve().is_relative_to(root) and path.is_file();iw,ih=Image.open(path).size;ratio=min(w/iw,h/ih);ww,hh=iw*ratio,ih*ratio;xx,yy=x+(w-ww)/2,y+(h-hh)/2
    slide.shapes.add_picture(str(path),Pt(xx),Pt(yy),width=Pt(ww),height=Pt(hh));pdf.drawImage(str(path),xx,H-yy-hh,width=ww,height=hh,mask='auto')
   elif kind=='text':
    size=item['size'];boldness=item.get('bold',False);lines=wrap(item['text'],size,w,boldness);leading=size*1.25
    assert len(lines)*leading<=h+1,(number,item['text'],len(lines)*leading,h)
    frame=slide.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h)).text_frame;frame.clear();frame.word_wrap=False;frame.vertical_anchor=MSO_ANCHOR.TOP
    frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=Pt(0)
    for i,line in enumerate(lines):
     p=frame.paragraphs[0] if i==0 else frame.add_paragraph();p.text=line;p.line_spacing=Pt(leading);p.space_before=p.space_after=Pt(0);p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[item.get('align','left')]
     for run in p.runs:
      run.font.name=FONT;run.font.size=Pt(size);run.font.bold=boldness;run.font.color.rgb=rgb(item.get('color','12243E'));rp=run._r.get_or_add_rPr();ea=OxmlElement('a:ea');ea.set('typeface',FONT);rp.append(ea)
     pdf.setFont('BriefBold' if boldness else 'BriefRegular',size);pdf.setFillColorRGB(*color(item.get('color','12243E')));tw=pdfmetrics.stringWidth(line,'BriefBold' if boldness else 'BriefRegular',size);xx=x+(w-tw)/2 if item.get('align')=='center' else x+w-tw if item.get('align')=='right' else x;pdf.drawString(xx,H-y-size*.9-i*leading,line)
    audit.append({'slide':number,'text':item['text'],'font_pt':size,'lines':len(lines),'needed_h':len(lines)*leading,'available_h':h,'bounds_ok':True})
   else:raise ValueError(kind)
  slide.notes_slide.notes_text_frame.text=spec.get('notes','')
  pdf.showPage()
 prs.save(str(pptx_path));pdf.save();return audit
